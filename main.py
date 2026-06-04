

#this downside code is working 


# import streamlit as st
# import tempfile
# import os
# from docx import Document
# import fitz  # PyMuPDF
# import pytesseract
# import cv2
# import numpy as np
# import speech_recognition as sr
# import ollama

# # File processing functions
# def extract_text_from_pdf(pdf_bytes) -> str:
#     with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
#         tmp.write(pdf_bytes)
#         tmp_path = tmp.name
#     doc = fitz.open(tmp_path)
#     full_text = [page.get_text() for page in doc]
#     doc.close()
#     os.remove(tmp_path)
#     return "\n".join(full_text)

# def extract_text_from_docx(docx_bytes) -> str:
#     with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
#         tmp.write(docx_bytes)
#         tmp_path = tmp.name
#     doc = Document(tmp_path)
#     text = "\n".join(p.text for p in doc.paragraphs)
#     os.remove(tmp_path)
#     return text

# def extract_text_from_audio(audio_bytes, audio_suffix) -> str:
#     with tempfile.NamedTemporaryFile(delete=False, suffix=audio_suffix) as tmp:
#         tmp.write(audio_bytes)
#         tmp_path = tmp.name
#     recognizer = sr.Recognizer()
#     with sr.AudioFile(tmp_path) as source:
#         audio = recognizer.record(source)
#     try:
#         text = recognizer.recognize_google(audio)
#     except Exception:
#         text = "[Could not transcribe audio]"
#     os.remove(tmp_path)
#     return text

# def process_image_to_text(image_bytes) -> str:
#     nparr = np.frombuffer(image_bytes, np.uint8)
#     img = cv2.imdecode(nparr, cv2.IMREAD_COLOR)
#     if img is None:
#         return "[Could not read image for OCR]"
#     text = pytesseract.image_to_string(img)
#     if not text.strip():
#         return "[No text detected in image]"
#     return text

# def process_video_to_text(video_bytes, video_suffix) -> str:
#     return "[Video text/scene extraction not implemented]"

# def chunk_text(text, max_length=1000):
#     return [text[i:i+max_length] for i in range(0, len(text), max_length)]

# def generate_ollama_answer(question, context):
#     prompt = f"Use the following context to answer the question in detail:\nContext:\n{context}\n\nQuestion:\n{question}\n\nAnswer:\n"
#     try:
#         response = ollama.chat(
#             model="tinyllama",
#             messages=[{"role": "user", "content": prompt}]
#         )
#         return response['message']['content']
#     except Exception as e:
#         return f"[Error generating answer: {str(e)}]"

# # Streamlit UI
# st.title("Multimodal RAG Chatbot with Ollama LLaMA")
# st.write("Upload PDF, Word, images, video, and audio, then ask a question. Answers generated locally with Ollama LLaMA.")

# uploaded_files = st.file_uploader("Upload files", type=[
#     'pdf', 'docx', 'doc', 'jpg', 'jpeg', 'png', 'gif', 'bmp',
#     'mp4', 'avi', 'mov', 'mkv', 'wav', 'mp3', 'flac'
# ], accept_multiple_files=True)

# question = st.text_area("Enter your question for all files", height=100)

# if st.button("Get Answers"):
#     if not uploaded_files:
#         st.warning("Please upload at least one file.")
#     elif not question.strip():
#         st.warning("Please enter a question.")
#     else:
#         for uploaded_file in uploaded_files:
#             filename = uploaded_file.name
#             ext = filename.split(".")[-1].lower()
#             file_bytes = uploaded_file.read()
#             content_text = ""

#             if ext == "pdf":
#                 content_text = extract_text_from_pdf(file_bytes)
#             elif ext in ("docx", "doc"):
#                 content_text = extract_text_from_docx(file_bytes)
#             elif ext in ("jpg", "jpeg", "png", "gif", "bmp"):
#                 content_text = process_image_to_text(file_bytes)
#             elif ext in ("wav", "mp3", "flac"):
#                 content_text = extract_text_from_audio(file_bytes, f".{ext}")
#             elif ext in ("mp4", "avi", "mov", "mkv"):
#                 content_text = process_video_to_text(file_bytes, f".{ext}")
#             else:
#                 content_text = "[Unsupported file type or unreadable.]"

#             st.markdown(f"### File: {filename}")
#             st.markdown("#### Extracted content (snippet):")
#             snippet = content_text[:1000] + ("..." if len(content_text) > 1000 else "")
#             st.text(snippet if snippet else "[No content extracted]")

#             with st.spinner(f"Generating answer for {filename} via Ollama LLaMA..."):
#                 full_chunks = chunk_text(content_text)
#                 context_window = "\n\n".join(full_chunks[:2])
#                 answer = generate_ollama_answer(question, context_window)
#             st.markdown("#### Answer:")
#             st.write(answer)
#             st.markdown("---")

# st.info("Tip: Ensure Ollama CLI is installed and configured, and LLaMA model is pulled locally.")


import streamlit as st
import tempfile
import os
from docx import Document
import fitz  # PyMuPDF
import pytesseract
import cv2
import numpy as np
import speech_recognition as sr
import ollama
from fpdf import FPDF

# Load and inject CSS style from external file
def load_css(filepath='styles.css'):
    with open(filepath) as f:
        st.markdown(f"<style>{f.read()}</style>", unsafe_allow_html=True)

load_css()  # Inject CSS at the start of app

# -------------------------------
# Helper functions (unchanged)
# -------------------------------

def extract_text_from_pdf(pdf_bytes) -> str:
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
        tmp.write(pdf_bytes)
        tmp_path = tmp.name
    doc = fitz.open(tmp_path)
    full_text = [page.get_text() for page in doc]
    doc.close()
    os.remove(tmp_path)
    return "\n".join(full_text)

def extract_text_from_docx(docx_bytes) -> str:
    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
        tmp.write(docx_bytes)
        tmp_path = tmp.name
    doc = Document(tmp_path)
    text = "\n".join(p.text for p in doc.paragraphs)
    os.remove(tmp_path)
    return text

def extract_text_from_audio(audio_bytes, audio_suffix) -> str:
    with tempfile.NamedTemporaryFile(delete=False, suffix=audio_suffix) as tmp:
        tmp.write(audio_bytes)
        tmp_path = tmp.name
    recognizer = sr.Recognizer()
    with sr.AudioFile(tmp_path) as source:
        audio = recognizer.record(source)
    try:
        text = recognizer.recognize_google(audio)
    except Exception:
        text = "[Could not transcribe audio]"
    os.remove(tmp_path)
    return text

def process_image_to_text(image_bytes) -> str:
    nparr = np.frombuffer(image_bytes, np.uint8)
    img = cv2.imdecode(nparr, cv2.IMREAD_COLOR)
    if img is None:
        return "[Could not read image for OCR]"
    text = pytesseract.image_to_string(img)
    if not text.strip():
        return "[No text detected in image]"
    return text

def process_video_to_text(video_bytes, video_suffix) -> str:
    return "[Video text/scene extraction not implemented]"

def chunk_text(text, max_length=1000):
    return [text[i:i+max_length] for i in range(0, len(text), max_length)]

def generate_ollama_answer(question, context):
    prompt = f"Use the following context to answer the question in detail:\nContext:\n{context}\n\nQuestion:\n{question}\n\nAnswer:\n"
    try:
        response = ollama.chat(
            model="tinyllama",
            messages=[{"role": "user", "content": prompt}]
        )
        return response['message']['content']
    except Exception as e:
        return f"[Error generating answer: {str(e)}]"

def clean_text_for_pdf(text):
    return text.encode('latin-1', 'replace').decode('latin-1')

def save_conversation_as_pdf(conversation, filename):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(0, 10, "Chatbot Conversation Transcript", ln=True, align="C")
    pdf.ln(10)
    for turn in conversation:
        role = clean_text_for_pdf(turn['role'])
        content = clean_text_for_pdf(turn['content'].replace('\n', ' '))
        pdf.multi_cell(0, 10, f"{role}: {content}")
        pdf.ln(2)
    pdf.output(filename)

# -------------------------------
# Streamlit UI
# -------------------------------

st.title("Multimodal Chatbot with Ollama with LLaMA")
st.write("Upload PDF, Word, image or audio — then ask a question. Answers generated locally using Ollama LLaMA.")

if 'transcript' not in st.session_state:
    st.session_state['transcript'] = []

uploaded_files = st.file_uploader(
    "Upload files",
    type=['pdf', 'docx', 'doc', 'jpg', 'jpeg', 'png', 'gif', 'bmp',
          'mp4', 'avi', 'mov', 'mkv', 'wav', 'mp3', 'flac'],
    accept_multiple_files=True
)

question = st.text_area("Enter your question for all files", height=100)

if st.button("Get Answers"):
    if not uploaded_files:
        st.warning("Please upload at least one file.")
    elif not question.strip():
        st.warning("Please enter a question.")
    else:
        for uploaded_file in uploaded_files:
            filename = uploaded_file.name
            ext = filename.split(".")[-1].lower()
            file_bytes = uploaded_file.read()
            content_text = ""

            if ext == "pdf":
                content_text = extract_text_from_pdf(file_bytes)
            elif ext in ("docx", "doc"):
                content_text = extract_text_from_docx(file_bytes)
            elif ext in ("jpg", "jpeg", "png", "gif", "bmp"):
                content_text = process_image_to_text(file_bytes)
            elif ext in ("wav", "mp3", "flac"):
                content_text = extract_text_from_audio(file_bytes, f".{ext}")
            elif ext in ("mp4", "avi", "mov", "mkv"):
                content_text = process_video_to_text(file_bytes, f".{ext}")
            else:
                content_text = "[Unsupported file type or unreadable.]"

            st.markdown(f"### File: {filename}")
            st.markdown("#### Extracted content (snippet):")
            snippet = content_text[:1000] + ("..." if len(content_text) > 1000 else "")
            st.text(snippet if snippet else "[No content extracted]")

            with st.spinner(f"Generating answer for {filename} via Ollama LLaMA..."):
                full_chunks = chunk_text(content_text)
                context_window = "\n\n".join(full_chunks[:2])
                answer = generate_ollama_answer(question, context_window)

            st.session_state.transcript.append({"role": "User", "content": question})
            st.session_state.transcript.append({"role": "Assistant", "content": answer})

            st.markdown("#### Answer:")
            st.write(answer)
            st.markdown("---")

if st.session_state.transcript:
    st.markdown("### Download Conversation Transcript")
    if st.button("Generate and Download PDF"):
        with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp_pdf:
            save_conversation_as_pdf(st.session_state.transcript, tmp_pdf.name)
            tmp_pdf.flush()
            with open(tmp_pdf.name, "rb") as f:
                pdf_data = f.read()
        st.download_button(
            label="Click to Download PDF",
            data=pdf_data,
            file_name="conversation.pdf",
            mime="application/pdf"
        )
